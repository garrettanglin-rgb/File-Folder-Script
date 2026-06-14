"""Generate a print-ready Avery label document from spreadsheet data.

Builds a new Word document from scratch, using exact dimensions extracted
from the Avery template profile (page size, margins, column widths, row
heights, cell margins).  This avoids any template-level baggage (document
protection, floating shapes, content controls) that can make the output
non-editable.
"""

import math
import platform
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips, RGBColor

from label_maker.config import (
    AVERY_TEMPLATE_PATH,
    FIELD_MAPPING,
    LABEL_LEFT_INDENT_PT,
    LINE_FORMAT_OVERRIDE,
    LINE_MAX_CHARS,
    NUMBERS_SPREADSHEET_PATH,
)
from label_maker.data_reader import clear_print_flags, read_flagged_rows
from label_maker.template_analyzer import load_profile

# Where the generated document is saved
OUTPUT_PATH = Path("~/Desktop/Labels_Ready.docx").expanduser()

# Where the cached formatting profile lives
PROFILE_PATH = AVERY_TEMPLATE_PATH.parent / "label_format.json"

# ---------------------------------------------------------------------------
# Alignment / spacing look-ups
# ---------------------------------------------------------------------------
_ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

_SPACING_RULE_MAP = {
    "single": WD_LINE_SPACING.SINGLE,
    "1.5": WD_LINE_SPACING.ONE_POINT_FIVE,
    "double": WD_LINE_SPACING.DOUBLE,
    "exactly": WD_LINE_SPACING.EXACTLY,
    "at_least": WD_LINE_SPACING.AT_LEAST,
    "multiple": WD_LINE_SPACING.MULTIPLE,
}


# ---------------------------------------------------------------------------
# Text formatting helpers
# ---------------------------------------------------------------------------

def _set_paragraph_spacing(paragraph, line_fmt):
    """Apply paragraph-level spacing and indent properties."""
    pf = paragraph.paragraph_format
    rule = _SPACING_RULE_MAP.get(line_fmt["line_spacing_rule"],
                                 WD_LINE_SPACING.SINGLE)
    pf.line_spacing_rule = rule
    if rule in (WD_LINE_SPACING.EXACTLY, WD_LINE_SPACING.AT_LEAST):
        pf.line_spacing = Pt(line_fmt["line_spacing_value"])
    else:
        pf.line_spacing = line_fmt["line_spacing_value"]

    sb = line_fmt.get("space_before_pt")
    sa = line_fmt.get("space_after_pt")
    pf.space_before = Pt(sb) if sb else Pt(0)
    pf.space_after = Pt(sa) if sa else Pt(0)

    # Left indent — pushes text past the colored tab area on the label.
    # Config value always wins; fall back to template-detected indent.
    li = LABEL_LEFT_INDENT_PT or line_fmt.get("left_indent_pt")
    if li:
        pf.left_indent = Pt(li)
    fli = line_fmt.get("first_line_indent_pt")
    if fli:
        pf.first_line_indent = Pt(fli)


def _format_run(run, line_fmt):
    """Apply run-level font properties from a line format dict."""
    font = run.font
    # Always set font name — fall back to Times New Roman if not detected
    font.name = line_fmt.get("font_name") or "Times New Roman"
    # Also set hAnsi so Word uses the same font for Latin characters
    r = run._r
    rPr = r.find(qn("w:rPr"))
    if rPr is not None:
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is not None:
            rFonts.set(qn("w:hAnsi"), font.name)
    # Always set font size — fall back to 12pt if not detected
    size = line_fmt.get("font_size_pt")
    font.size = Pt(size) if size else Pt(12)
    font.bold = line_fmt.get("bold", False)
    font.italic = line_fmt.get("italic", False)
    font.underline = line_fmt.get("underline", False)
    if line_fmt.get("color_rgb"):
        font.color.rgb = RGBColor.from_string(line_fmt["color_rgb"])


def _add_right_tab_stop(paragraph, position_twips):
    """Add a right-aligned tab stop to a paragraph via XML."""
    pPr = paragraph._p.get_or_add_pPr()
    tabs = pPr.find(qn("w:tabs"))
    if tabs is None:
        tabs = OxmlElement("w:tabs")
        pPr.append(tabs)
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), str(position_twips))
    tab.set(qn("w:leader"), "none")
    tabs.append(tab)


def _apply_line_single(paragraph, line_fmt, text):
    """Fill a paragraph with a single column value."""
    paragraph.alignment = _ALIGN_MAP.get(line_fmt["alignment"],
                                         WD_ALIGN_PARAGRAPH.LEFT)
    _set_paragraph_spacing(paragraph, line_fmt)
    run = paragraph.add_run(str(text) if text is not None else "")
    _format_run(run, line_fmt)


def _apply_line_pair(paragraph, line_fmt, left_text, right_text,
                     cell_width_twips, cell_margin_lr_twips):
    """Fill a paragraph with two values: left-aligned and right-aligned.

    Uses a right-aligned tab stop so the second value sits flush-right
    inside the cell.
    """
    # Override to left-align so the tab stop layout works correctly
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_paragraph_spacing(paragraph, line_fmt)

    # Tab stop position = cell content width minus the left indent,
    # so the right-aligned text doesn't go past the label edge.
    indent_twips = int((LABEL_LEFT_INDENT_PT or 0) * 20)  # 1pt = 20 twips
    tab_pos = cell_width_twips - (2 * cell_margin_lr_twips) - indent_twips
    _add_right_tab_stop(paragraph, tab_pos)

    # Left run
    run_left = paragraph.add_run(
        str(left_text) if left_text is not None else "")
    _format_run(run_left, line_fmt)

    # Tab character
    run_tab = paragraph.add_run("\t")
    _format_run(run_tab, line_fmt)

    # Right run
    run_right = paragraph.add_run(
        str(right_text) if right_text is not None else "")
    _format_run(run_right, line_fmt)


def _truncate(text, max_chars):
    """Truncate text to max_chars, adding '...' if shortened."""
    text = str(text) if text is not None else ""
    if max_chars and len(text) > max_chars:
        return text[:max_chars - 3].rstrip() + "..."
    return text


def _get_all_mapped_columns(field_mapping):
    """Return a flat set of every column name referenced in FIELD_MAPPING."""
    cols = set()
    for value in field_mapping.values():
        if isinstance(value, tuple):
            cols.update(value)
        else:
            cols.add(value)
    return cols


def _populate_cell(cell, data_row, line_formats, field_mapping,
                   cell_width_twips, cell_margin_lr_twips):
    """Fill a single table cell with label data using the format profile.

    *field_mapping* values may be a string (single column → full line)
    or a tuple of two strings (left column, right column on one line).
    """
    # Fallback format if the template had no detectable line styles
    _DEFAULT_FMT = {
        "font_name": "Times New Roman", "font_size_pt": 12.0,
        "bold": False, "italic": False, "underline": False,
        "alignment": "left", "line_spacing_rule": "single",
        "line_spacing_value": 1.0, "space_before_pt": 0.0,
        "space_after_pt": 0.0, "color_rgb": None,
    }

    sorted_lines = sorted(field_mapping.items())

    for li, (line_num, columns) in enumerate(sorted_lines):
        if line_formats:
            fmt_index = min(li, len(line_formats) - 1)
            fmt = dict(line_formats[fmt_index])  # copy so we can override
        else:
            fmt = dict(_DEFAULT_FMT)

        # Apply per-line formatting overrides from config
        overrides = LINE_FORMAT_OVERRIDE.get(line_num, {})
        fmt.update(overrides)

        para = cell.paragraphs[0] if li == 0 else cell.add_paragraph()

        # Truncate if a max character limit is set for this line
        max_chars = LINE_MAX_CHARS.get(line_num)

        if isinstance(columns, tuple):
            left_col, right_col = columns
            _apply_line_pair(
                para, fmt,
                _truncate(data_row.get(left_col, ""), max_chars),
                _truncate(data_row.get(right_col, ""), max_chars),
                cell_width_twips,
                cell_margin_lr_twips,
            )
        else:
            _apply_line_single(
                para, fmt,
                _truncate(data_row.get(columns, ""), max_chars),
            )


# ---------------------------------------------------------------------------
# Build document from scratch using profile dimensions
# ---------------------------------------------------------------------------

def _build_page_table(doc, profile):
    """Create a table matching the Avery template's exact dimensions.

    Reads column widths, row heights, and cell margins from the profile
    and applies them via XML so the printed output aligns with the
    physical label sheet.

    Returns (table, label_row_indices, label_col_indices).
    """
    grid = profile["grid"]
    cell_info = profile.get("cell", {})
    num_rows = grid["num_rows"]
    num_cols = grid["num_cols"]
    label_row_indices = grid.get("label_row_indices", list(range(num_rows)))
    label_col_indices = grid.get("label_col_indices", list(range(num_cols)))

    col_widths = cell_info.get("col_widths_twips", [])
    all_row_heights = cell_info.get("all_row_heights_twips",
                                     cell_info.get("row_heights_twips", []))
    margin_top = cell_info.get("margin_top_twips")
    margin_bottom = cell_info.get("margin_bottom_twips")
    margin_left = cell_info.get("margin_left_twips")
    margin_right = cell_info.get("margin_right_twips")

    # Create the table
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.autofit = False
    tbl = table._tbl

    # Remove default borders (Avery labels typically have no visible borders)
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)

    # Set table width to the sum of column widths
    total_width = sum(col_widths) if col_widths else num_cols * 5000
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(total_width))
    tblW.set(qn("w:type"), "dxa")

    # Set table layout to fixed
    tblLayout = OxmlElement("w:tblLayout")
    tblLayout.set(qn("w:type"), "fixed")
    tblPr.append(tblLayout)

    # Set table-level cell margins
    if any(v is not None for v in [margin_top, margin_bottom,
                                    margin_left, margin_right]):
        tblCellMar = OxmlElement("w:tblCellMar")
        for side, val in [("top", margin_top), ("bottom", margin_bottom),
                          ("left", margin_left), ("right", margin_right)]:
            if val is not None:
                el = OxmlElement(f"w:{side}")
                el.set(qn("w:w"), str(val))
                el.set(qn("w:type"), "dxa")
                tblCellMar.append(el)
        tblPr.append(tblCellMar)

    # Remove all borders
    tblBorders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "none")
        border.set(qn("w:sz"), "0")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "auto")
        tblBorders.append(border)
    tblPr.append(tblBorders)

    # Set column widths via tblGrid
    tblGrid = tbl.find(qn("w:tblGrid"))
    if tblGrid is not None:
        tbl.remove(tblGrid)
    tblGrid = OxmlElement("w:tblGrid")
    for ci in range(num_cols):
        gridCol = OxmlElement("w:gridCol")
        w = col_widths[ci] if ci < len(col_widths) else 5000
        gridCol.set(qn("w:w"), str(w))
        tblGrid.append(gridCol)
    # Insert tblGrid right after tblPr
    tblPr_idx = list(tbl).index(tblPr)
    tbl.insert(tblPr_idx + 1, tblGrid)

    # Set row heights and cell widths
    label_row_set = set(label_row_indices)
    label_col_set = set(label_col_indices)

    for ri, row in enumerate(table.rows):
        tr = row._tr
        # Set row height
        if ri < len(all_row_heights) and all_row_heights[ri] is not None:
            trPr = tr.find(qn("w:trPr"))
            if trPr is None:
                trPr = OxmlElement("w:trPr")
                tr.insert(0, trPr)
            trHeight = OxmlElement("w:trHeight")
            trHeight.set(qn("w:val"), str(all_row_heights[ri]))
            trHeight.set(qn("w:hRule"), "exact")
            trPr.append(trHeight)

        for ci, cell in enumerate(row.cells):
            tc = cell._tc
            # Set cell width
            tcPr = tc.find(qn("w:tcPr"))
            if tcPr is None:
                tcPr = OxmlElement("w:tcPr")
                tc.insert(0, tcPr)
            tcW = OxmlElement("w:tcW")
            w = col_widths[ci] if ci < len(col_widths) else 5000
            tcW.set(qn("w:w"), str(w))
            tcW.set(qn("w:type"), "dxa")
            tcPr.append(tcW)

            # Label cells: pin text to top, zero out top margin
            if ri in label_row_set and ci in label_col_set:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
                tcMar = OxmlElement("w:tcMar")
                top_el = OxmlElement("w:top")
                top_el.set(qn("w:w"), "0")
                top_el.set(qn("w:type"), "dxa")
                tcMar.append(top_el)
                tcPr.append(tcMar)

    return table, label_row_indices, label_col_indices


def _fill_page_cells(table, data_rows, row_cursor,
                     label_row_indices, label_col_indices,
                     line_formats, field_mapping,
                     label_col_width, cell_margin_lr):
    """Fill label cells on one page, returning the updated row cursor."""
    for ri in label_row_indices:
        for ci in label_col_indices:
            if row_cursor < len(data_rows):
                _populate_cell(
                    table.rows[ri].cells[ci],
                    data_rows[row_cursor],
                    line_formats,
                    field_mapping,
                    label_col_width,
                    cell_margin_lr,
                )
                row_cursor += 1
    return row_cursor


def generate(profile_path=None, spreadsheet_path=None, output_path=None,
             field_mapping=None):
    """Run the full label-generation pipeline.

    Builds a new document from scratch using exact dimensions from the
    Avery template profile.  This produces a clean, fully editable .docx
    without any template-level baggage.

    Parameters
    ----------
    profile_path : str or Path, optional
        Path to label_format.json.  Defaults to ``PROFILE_PATH``.
    spreadsheet_path : str or Path, optional
        Path to the .numbers file.  Defaults to ``NUMBERS_SPREADSHEET_PATH``.
    output_path : str or Path, optional
        Where to save the generated .docx.  Defaults to ``OUTPUT_PATH``.
    field_mapping : dict, optional
        ``{line_number: column_name}`` mapping.  Defaults to
        ``config.FIELD_MAPPING``.

    Returns
    -------
    Path or None
        The path to the saved document, or *None* if there was nothing
        to print.
    """
    profile_path = Path(profile_path or PROFILE_PATH)
    spreadsheet_path = Path(spreadsheet_path or NUMBERS_SPREADSHEET_PATH)
    output_path = Path(output_path or OUTPUT_PATH)
    if field_mapping is None:
        field_mapping = FIELD_MAPPING

    # 1. Load the formatting profile (re-analyze if template is newer)
    from label_maker.template_analyzer import analyze_template, save_profile
    if (not profile_path.exists()
            or AVERY_TEMPLATE_PATH.stat().st_mtime > profile_path.stat().st_mtime):
        print("Analyzing template (new or updated)...")
        profile = analyze_template(AVERY_TEMPLATE_PATH)
        save_profile(profile, profile_path)
    else:
        profile = load_profile(profile_path)
    grid = profile["grid"]
    fmt = profile["format_template"]
    labels_per_page = grid["labels_per_page"]

    # 2. Read flagged rows
    headers, rows, indices = read_flagged_rows(spreadsheet_path)

    if not rows:
        print("No new labels to print.")
        return None

    # Validate that every mapped column exists in the spreadsheet
    mapped_cols = _get_all_mapped_columns(field_mapping)
    missing = mapped_cols - set(headers)
    if missing:
        raise ValueError(
            f"FIELD_MAPPING references columns not found in the spreadsheet: "
            f"{sorted(missing)}.  Available headers: {headers}"
        )

    print(f"Found {len(rows)} label(s) to print.")

    # 3. Create a new document with the template's page dimensions
    doc = Document()
    page = profile.get("page", {})
    section = doc.sections[0]

    if page.get("page_width_in"):
        section.page_width = Inches(page["page_width_in"])
    if page.get("page_height_in"):
        section.page_height = Inches(page["page_height_in"])
    if page.get("margin_top_in"):
        section.top_margin = Inches(page["margin_top_in"])
    if page.get("margin_bottom_in"):
        section.bottom_margin = Inches(page["margin_bottom_in"])
    if page.get("margin_left_in"):
        section.left_margin = Inches(page["margin_left_in"])
    if page.get("margin_right_in"):
        section.right_margin = Inches(page["margin_right_in"])

    # Remove the default empty paragraph that Document() creates
    body = doc.element.body
    for p in body.findall(qn("w:p")):
        body.remove(p)

    # 4. Build pages with data
    cell_info = profile.get("cell", {})
    col_widths = cell_info.get("col_widths_twips", [])
    label_col_indices = grid.get(
        "label_col_indices", list(range(grid["num_cols"])))
    label_col_width = (
        col_widths[label_col_indices[0]]
        if col_widths and label_col_indices
        and label_col_indices[0] < len(col_widths)
        else 5000
    )
    cell_margin_lr = cell_info.get("margin_left_twips") or 115

    total_pages = math.ceil(len(rows) / labels_per_page)
    row_cursor = 0

    for page_num in range(total_pages):
        # Add a page break before every page except the first
        if page_num > 0:
            p_el = OxmlElement("w:p")
            r_el = OxmlElement("w:r")
            br_el = OxmlElement("w:br")
            br_el.set(qn("w:type"), "page")
            r_el.append(br_el)
            p_el.append(r_el)
            body.append(p_el)

        table, label_row_idx, label_col_idx = _build_page_table(doc, profile)
        row_cursor = _fill_page_cells(
            table, rows, row_cursor,
            label_row_idx, label_col_idx,
            fmt["line_formats"], field_mapping,
            label_col_width, cell_margin_lr,
        )

    # 5. Word requires a paragraph after every table.  Without one it
    #    auto-creates a default-sized paragraph that pushes to a blank
    #    second page.  Insert a near-invisible one before the sectPr.
    trailing_p = OxmlElement("w:p")
    trailing_pPr = OxmlElement("w:pPr")
    # Tiny font so the paragraph mark takes virtually no space
    rPr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "2")          # 1pt font
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), "2")
    rPr.append(sz)
    rPr.append(szCs)
    trailing_pPr.append(rPr)
    # Zero spacing, minimal line height
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "20")   # 1pt line height
    spacing.set(qn("w:lineRule"), "exact")
    trailing_pPr.append(spacing)
    trailing_p.append(trailing_pPr)
    # Insert before sectPr so it sits right after the last table
    sect_pr = body.find(qn("w:sectPr"))
    if sect_pr is not None:
        sect_pr.addprevious(trailing_p)
    else:
        body.append(trailing_p)

    # 6. Save the document
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"Saved {len(rows)} label(s) across {total_pages} page(s) "
          f"to {output_path}")

    # 6. Clear the print flags so they don't print again
    clear_print_flags(spreadsheet_path, indices)

    # 7. Open the document
    _open_file(output_path)

    return output_path


def _open_file(path):
    """Open a file with the system default application."""
    path = str(path)
    system = platform.system()

    # On macOS, remove the quarantine extended attribute that makes Word
    # open downloaded files in Protected View (read-only).
    if system == "Darwin":
        try:
            subprocess.run(
                ["xattr", "-d", "com.apple.quarantine", path],
                capture_output=True,
            )
        except OSError:
            pass

    try:
        if system == "Darwin":
            subprocess.Popen(["open", path])
        elif system == "Windows":
            subprocess.Popen(["start", "", path], shell=True)
        else:
            subprocess.Popen(["xdg-open", path])
    except OSError:
        print(f"Could not auto-open {path}. Please open it manually.")


# ---------------------------------------------------------------------------
# CLI entry point
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    generate()
