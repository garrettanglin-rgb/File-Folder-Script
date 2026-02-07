"""Analyze an Avery label Word document and extract its formatting profile."""

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn


# Map python-docx alignment enums to readable strings
_ALIGNMENT_MAP = {
    WD_ALIGN_PARAGRAPH.LEFT: "left",
    WD_ALIGN_PARAGRAPH.CENTER: "center",
    WD_ALIGN_PARAGRAPH.RIGHT: "right",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
    None: "left",  # Word default when unset
}

# Map line-spacing rule enums to readable strings
_SPACING_RULE_MAP = {
    WD_LINE_SPACING.SINGLE: "single",
    WD_LINE_SPACING.ONE_POINT_FIVE: "1.5",
    WD_LINE_SPACING.DOUBLE: "double",
    WD_LINE_SPACING.EXACTLY: "exactly",
    WD_LINE_SPACING.AT_LEAST: "at_least",
    WD_LINE_SPACING.MULTIPLE: "multiple",
    None: "single",
}


def _emu_to_pt(emu_val):
    """Convert EMU (English Metric Units) to points. Returns None if input is None."""
    if emu_val is None:
        return None
    return round(emu_val / 12700, 1)


def _emu_to_inches(emu_val):
    """Convert EMU to inches. Returns None if input is None."""
    if emu_val is None:
        return None
    return round(emu_val / 914400, 4)


def _get_document_defaults(doc):
    """Extract document-level default font properties from w:docDefaults."""
    defaults = {
        "font_name": None,
        "font_size_pt": None,
        "bold": False,
        "italic": False,
        "underline": False,
        "color_rgb": None,
    }

    styles_element = doc.styles.element
    doc_defaults_el = styles_element.find(qn("w:docDefaults"))
    if doc_defaults_el is None:
        return defaults

    rpr_default = doc_defaults_el.find(qn("w:rPrDefault"))
    if rpr_default is None:
        return defaults

    rpr = rpr_default.find(qn("w:rPr"))
    if rpr is None:
        return defaults

    # Font name from rFonts
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is not None:
        defaults["font_name"] = (
            rFonts.get(qn("w:ascii"))
            or rFonts.get(qn("w:hAnsi"))
            or rFonts.get(qn("w:cs"))
        )

    # Font size (stored as half-points in XML)
    sz = rpr.find(qn("w:sz"))
    if sz is not None:
        half_points = sz.get(qn("w:val"))
        if half_points:
            defaults["font_size_pt"] = int(half_points) / 2.0

    # Bold
    b_el = rpr.find(qn("w:b"))
    if b_el is not None:
        val = b_el.get(qn("w:val"))
        defaults["bold"] = val is None or val.lower() in ("true", "1", "on")

    # Italic
    i_el = rpr.find(qn("w:i"))
    if i_el is not None:
        val = i_el.get(qn("w:val"))
        defaults["italic"] = val is None or val.lower() in ("true", "1", "on")

    # Underline
    u_el = rpr.find(qn("w:u"))
    if u_el is not None:
        val = u_el.get(qn("w:val"))
        defaults["underline"] = val is not None and val != "none"

    return defaults


def _extract_cell_dimensions(table):
    """Extract column widths, row heights, and cell margins from the table XML."""
    tbl = table._tbl

    # Column widths from w:tblGrid/w:gridCol
    col_widths = []
    tbl_grid = tbl.find(qn("w:tblGrid"))
    if tbl_grid is not None:
        for grid_col in tbl_grid.findall(qn("w:gridCol")):
            w = grid_col.get(qn("w:w"))
            if w:
                col_widths.append(int(w))

    # Row heights from w:trPr/w:trHeight
    row_heights = []
    for tr in tbl.findall(qn("w:tr")):
        height = None
        trPr = tr.find(qn("w:trPr"))
        if trPr is not None:
            trHeight = trPr.find(qn("w:trHeight"))
            if trHeight is not None:
                val = trHeight.get(qn("w:val"))
                if val:
                    height = int(val)
        row_heights.append(height)

    # Cell margins — first from individual cell (w:tc/w:tcPr/w:tcMar)
    margins = {"top": None, "bottom": None, "left": None, "right": None}
    first_row_trs = tbl.findall(qn("w:tr"))
    if first_row_trs:
        first_tcs = first_row_trs[0].findall(qn("w:tc"))
        if first_tcs:
            tcPr = first_tcs[0].find(qn("w:tcPr"))
            if tcPr is not None:
                tcMar = tcPr.find(qn("w:tcMar"))
                if tcMar is not None:
                    for side in ("top", "bottom", "left", "right"):
                        el = tcMar.find(qn(f"w:{side}"))
                        if el is not None:
                            w = el.get(qn("w:w"))
                            if w:
                                margins[side] = int(w)

    # Fall back to table-level default cell margins (w:tblPr/w:tblCellMar)
    if all(v is None for v in margins.values()):
        tblPr = tbl.find(qn("w:tblPr"))
        if tblPr is not None:
            tblCellMar = tblPr.find(qn("w:tblCellMar"))
            if tblCellMar is not None:
                for side in ("top", "bottom", "left", "right"):
                    el = tblCellMar.find(qn(f"w:{side}"))
                    if el is not None:
                        w = el.get(qn("w:w"))
                        if w:
                            margins[side] = int(w)

    return col_widths, row_heights, margins


def _detect_label_grid(col_widths, row_heights):
    """Identify which columns/rows are actual label cells vs gutters/spacers.

    Gutter columns are significantly narrower than label columns.
    Spacer rows are significantly shorter than label rows.
    A column/row whose size is less than 40% of the largest is treated as a
    gutter/spacer.

    Returns (label_col_indices, label_row_indices).
    """
    # Detect gutter columns
    if col_widths:
        max_w = max(col_widths)
        label_cols = [i for i, w in enumerate(col_widths) if w >= max_w * 0.4]
    else:
        label_cols = []

    # Detect spacer rows
    valid_heights = [h for h in row_heights if h is not None]
    if valid_heights:
        max_h = max(valid_heights)
        label_rows = []
        for i, h in enumerate(row_heights):
            if h is None or h >= max_h * 0.4:
                label_rows.append(i)
    else:
        label_rows = list(range(len(row_heights)))

    return label_cols, label_rows


def _analyze_run(run, para_style, doc_defaults):
    """Extract formatting from a single text run, resolving via style hierarchy."""
    font = run.font

    # Run-level values
    name = font.name
    size_pt = _emu_to_pt(font.size)
    bold = font.bold
    italic = font.italic
    underline = font.underline
    color = str(font.color.rgb) if font.color and font.color.rgb else None

    # Resolve None values: run → paragraph style → document defaults
    style_font = para_style.font if para_style else None

    if name is None and style_font and style_font.name:
        name = style_font.name
    if name is None and doc_defaults:
        name = doc_defaults.get("font_name")

    if size_pt is None and style_font and style_font.size:
        size_pt = _emu_to_pt(style_font.size)
    if size_pt is None and doc_defaults:
        size_pt = doc_defaults.get("font_size_pt")

    if bold is None and style_font:
        bold = style_font.bold
    if bold is None and doc_defaults:
        bold = doc_defaults.get("bold", False)
    bold = bool(bold) if bold is not None else False

    if italic is None and style_font:
        italic = style_font.italic
    if italic is None and doc_defaults:
        italic = doc_defaults.get("italic", False)
    italic = bool(italic) if italic is not None else False

    if underline is None and style_font:
        underline = style_font.underline
    if underline is None and doc_defaults:
        underline = doc_defaults.get("underline", False)
    underline = bool(underline) if underline is not None else False

    if color is None and style_font:
        try:
            if style_font.color and style_font.color.rgb:
                color = str(style_font.color.rgb)
        except Exception:
            pass
    if color is None and doc_defaults:
        color = doc_defaults.get("color_rgb")

    return {
        "text": run.text,
        "font_name": name,
        "font_size_pt": size_pt,
        "bold": bold,
        "italic": italic,
        "underline": underline,
        "color_rgb": color,
    }


def _analyze_paragraph(para, doc_defaults):
    """Extract formatting from a single paragraph (one text line in a label)."""
    pf = para.paragraph_format
    alignment = _ALIGNMENT_MAP.get(para.alignment, "left")

    # Determine line spacing value and rule
    spacing_rule = _SPACING_RULE_MAP.get(pf.line_spacing_rule, "single")
    if pf.line_spacing is None:
        line_spacing_value = 1.0
    elif pf.line_spacing_rule in (WD_LINE_SPACING.EXACTLY, WD_LINE_SPACING.AT_LEAST):
        # Value is in EMU — convert to points
        line_spacing_value = _emu_to_pt(pf.line_spacing)
    else:
        # Value is a float multiplier (e.g. 1.0, 1.5, 2.0)
        line_spacing_value = float(pf.line_spacing)

    para_style = para.style
    runs = [_analyze_run(r, para_style, doc_defaults) for r in para.runs]

    # Derive dominant / first-run formatting for convenience.
    if runs:
        first = runs[0]
    else:
        # No runs (blank template cell) — resolve from style → doc defaults
        style_font = para_style.font if para_style else None
        name = None
        size_pt = None
        bold = False
        italic = False
        underline_val = False
        color = None

        if style_font:
            name = style_font.name
            size_pt = _emu_to_pt(style_font.size) if style_font.size else None
            bold = bool(style_font.bold) if style_font.bold is not None else False
            italic = bool(style_font.italic) if style_font.italic is not None else False
            underline_val = bool(style_font.underline) if style_font.underline is not None else False
            try:
                if style_font.color and style_font.color.rgb:
                    color = str(style_font.color.rgb)
            except Exception:
                pass

        # Fall back to document defaults for any remaining Nones
        if name is None and doc_defaults:
            name = doc_defaults.get("font_name")
        if size_pt is None and doc_defaults:
            size_pt = doc_defaults.get("font_size_pt")
        if not bold and doc_defaults:
            bold = doc_defaults.get("bold", False)
        if not italic and doc_defaults:
            italic = doc_defaults.get("italic", False)
        if not underline_val and doc_defaults:
            underline_val = doc_defaults.get("underline", False)
        if color is None and doc_defaults:
            color = doc_defaults.get("color_rgb")

        first = {
            "font_name": name,
            "font_size_pt": size_pt,
            "bold": bold,
            "italic": italic,
            "underline": underline_val,
            "color_rgb": color,
        }

    return {
        "text": para.text,
        "alignment": alignment,
        "line_spacing_rule": spacing_rule,
        "line_spacing_value": line_spacing_value,
        "space_before_pt": _emu_to_pt(pf.space_before),
        "space_after_pt": _emu_to_pt(pf.space_after),
        "left_indent_pt": _emu_to_pt(pf.left_indent),
        "first_line_indent_pt": _emu_to_pt(pf.first_line_indent),
        "font_name": first.get("font_name"),
        "font_size_pt": first.get("font_size_pt"),
        "bold": first.get("bold", False),
        "italic": first.get("italic", False),
        "underline": first.get("underline", False),
        "color_rgb": first.get("color_rgb"),
        "runs": runs,
    }


def _analyze_cell(cell, doc_defaults):
    """Analyze all paragraphs (lines) inside a single label cell.

    If the cell has text, only paragraphs with visible text are kept.
    If the cell is blank (clean template), *all* paragraphs are analyzed
    so we still capture font/size/alignment from the empty runs.
    """
    paragraphs = cell.paragraphs
    lines_with_text = [_analyze_paragraph(p, doc_defaults)
                       for p in paragraphs if p.text.strip()]
    if lines_with_text:
        return {"num_lines": len(lines_with_text), "lines": lines_with_text}

    # Blank cell — analyze every paragraph for its formatting
    all_lines = [_analyze_paragraph(p, doc_defaults) for p in paragraphs]
    return {"num_lines": len(all_lines), "lines": all_lines}


def analyze_template(docx_path):
    """Open an Avery label .docx and return a complete formatting profile.

    Parameters
    ----------
    docx_path : str or Path
        Path to the Word document to analyze.

    Returns
    -------
    dict
        A dictionary describing the page layout, grid structure, cell
        dimensions, and per-line formatting of every label in the template.
    """
    docx_path = Path(docx_path)
    doc = Document(str(docx_path))

    # --- Document defaults (font inheritance) ---
    doc_defaults = _get_document_defaults(doc)

    # --- Page / section info ---
    section = doc.sections[0]
    page_info = {
        "page_width_in": _emu_to_inches(section.page_width),
        "page_height_in": _emu_to_inches(section.page_height),
        "margin_top_in": _emu_to_inches(section.top_margin),
        "margin_bottom_in": _emu_to_inches(section.bottom_margin),
        "margin_left_in": _emu_to_inches(section.left_margin),
        "margin_right_in": _emu_to_inches(section.right_margin),
    }

    # --- Table / grid detection ---
    if not doc.tables:
        raise ValueError(
            "No tables found in the document. "
            "Avery label templates use a table grid."
        )

    table = doc.tables[0]
    num_rows = len(table.rows)
    num_cols = len(table.columns)

    # --- Cell dimensions from XML ---
    col_widths, row_heights, margins = _extract_cell_dimensions(table)

    # --- Detect gutter columns and spacer rows ---
    label_col_indices, label_row_indices = _detect_label_grid(
        col_widths, row_heights
    )
    if not label_col_indices:
        label_col_indices = list(range(num_cols))
    if not label_row_indices:
        label_row_indices = list(range(num_rows))

    labels_per_page = len(label_row_indices) * len(label_col_indices)

    grid_info = {
        "num_rows": num_rows,
        "num_cols": num_cols,
        "label_rows": len(label_row_indices),
        "label_cols": len(label_col_indices),
        "label_row_indices": label_row_indices,
        "label_col_indices": label_col_indices,
        "labels_per_page": labels_per_page,
    }

    cell_info = {
        "col_widths_twips": col_widths,
        "row_heights_twips": [h for h in row_heights if h is not None],
        "all_row_heights_twips": row_heights,
        "margin_top_twips": margins["top"],
        "margin_bottom_twips": margins["bottom"],
        "margin_left_twips": margins["left"],
        "margin_right_twips": margins["right"],
    }

    # --- Per-label analysis (only actual label cells, not gutters/spacers) ---
    labels = []
    label_col_set = set(label_col_indices)
    label_row_set = set(label_row_indices)
    for ri, row in enumerate(table.rows):
        if ri not in label_row_set:
            continue
        for ci, cell in enumerate(row.cells):
            if ci not in label_col_set:
                continue
            label_data = _analyze_cell(cell, doc_defaults)
            label_data["row"] = ri
            label_data["col"] = ci
            labels.append(label_data)

    # --- Build a canonical "format template" from the first label ---
    first_label = labels[0] if labels else {}
    line_formats = []
    for line in first_label.get("lines", []):
        line_formats.append({
            "font_name": line["font_name"],
            "font_size_pt": line["font_size_pt"],
            "bold": line["bold"],
            "italic": line["italic"],
            "underline": line["underline"],
            "alignment": line["alignment"],
            "line_spacing_rule": line["line_spacing_rule"],
            "line_spacing_value": line["line_spacing_value"],
            "space_before_pt": line["space_before_pt"],
            "space_after_pt": line["space_after_pt"],
            "left_indent_pt": line.get("left_indent_pt"),
            "first_line_indent_pt": line.get("first_line_indent_pt"),
            "color_rgb": line["color_rgb"],
        })

    format_template = {
        "num_lines_per_label": first_label.get("num_lines", 0),
        "line_formats": line_formats,
    }

    profile = {
        "source_file": str(docx_path.resolve()),
        "doc_defaults": doc_defaults,
        "page": page_info,
        "grid": grid_info,
        "cell": cell_info,
        "format_template": format_template,
        "labels": labels,
    }

    return profile


def save_profile(profile, output_path):
    """Write a formatting profile to a JSON file.

    Parameters
    ----------
    profile : dict
        The profile dictionary returned by ``analyze_template``.
    output_path : str or Path
        Destination path for the JSON file.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "w") as f:
        json.dump(profile, f, indent=2)
    print(f"Formatting profile saved to {output_path}")


def load_profile(json_path):
    """Load a previously saved formatting profile from JSON.

    Parameters
    ----------
    json_path : str or Path
        Path to the JSON profile file.

    Returns
    -------
    dict
        The formatting profile dictionary.
    """
    with open(json_path) as f:
        return json.load(f)


# ---------------------------------------------------------------------------
# CLI entry point
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    from label_maker.config import AVERY_TEMPLATE_PATH

    docx_file = Path(AVERY_TEMPLATE_PATH)
    output_file = docx_file.parent / "label_format.json"

    print(f"Analyzing template: {docx_file}")
    profile = analyze_template(docx_file)

    print(f"\nPage: {profile['page']['page_width_in']}\" x "
          f"{profile['page']['page_height_in']}\"")
    grid = profile["grid"]
    print(f"Physical table: {grid['num_rows']} rows x {grid['num_cols']} cols")
    print(f"Label grid: {grid['label_rows']} rows x {grid['label_cols']} cols "
          f"= {grid['labels_per_page']} labels/page")
    print(f"Lines per label: "
          f"{profile['format_template']['num_lines_per_label']}")

    cell = profile.get("cell", {})
    if cell.get("col_widths_twips"):
        print(f"Column widths (twips): {cell['col_widths_twips']}")
    if cell.get("row_heights_twips"):
        heights = cell["row_heights_twips"]
        preview = heights[:5]
        suffix = "..." if len(heights) > 5 else ""
        print(f"Row heights (twips): {preview}{suffix}")
    if cell.get("margin_left_twips") is not None:
        print(f"Cell margins L/R/T/B: "
              f"{cell.get('margin_left_twips')}/"
              f"{cell.get('margin_right_twips')}/"
              f"{cell.get('margin_top_twips')}/"
              f"{cell.get('margin_bottom_twips')}")

    defaults = profile.get("doc_defaults", {})
    if defaults.get("font_name") or defaults.get("font_size_pt"):
        print(f"Document defaults: {defaults.get('font_name')} "
              f"{defaults.get('font_size_pt')}pt")

    for i, lf in enumerate(profile["format_template"]["line_formats"]):
        style_parts = []
        if lf["bold"]:
            style_parts.append("Bold")
        if lf["italic"]:
            style_parts.append("Italic")
        if lf["underline"]:
            style_parts.append("Underline")
        style_str = "+".join(style_parts) if style_parts else "Regular"
        indent = lf.get('left_indent_pt')
        indent_str = f", indent={indent}pt" if indent else ""
        print(f"  Line {i + 1}: {lf['font_name']} {lf['font_size_pt']}pt, "
              f"{style_str}, {lf['alignment']}, spacing={lf['line_spacing_value']}"
              f"{indent_str}")

    save_profile(profile, output_file)
